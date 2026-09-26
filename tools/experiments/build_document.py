"""One synthetic entity-definition export; no general Plugin converter."""
from dataclasses import dataclass, field, fields
from pathlib import Path
from typing import Optional, get_type_hints
import argparse
import hashlib
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn


@dataclass
class Invoice:
    invoice_id: int = field(metadata=dict(pk=True, description="請求書の識別子", constraint="REQ-SYN-01"))
    customer_name: str = field(metadata=dict(pk=False, description="請求先の日本語名称", constraint="REQ-SYN-02"))
    note: Optional[str] = field(default=None, metadata=dict(pk=False,
        description="長文の確認欄。移行時に旧システムの備考を保持し、複数行へ折り返しても文字が欠落しないことを確認する。これは架空の説明であり、顧客情報は含まない。", constraint=""))


def main(out):
    out.mkdir(parents=True, exist_ok=True)
    hints = get_type_hints(Invoice)
    rows = []
    for i, f in enumerate(fields(Invoice), 1):
        t = hints[f.name]
        optional = t == Optional[str]
        rows.append([i, f.name, "Optional[str]" if optional else t.__name__, "○" if f.metadata["pk"] else "",
                     "" if optional else "○", f.metadata["description"], f.metadata["constraint"]])
    data = dict(title="エンティティ定義", entity="Invoice（架空の請求書）", headers=["#", "属性", "型", "PK", "必須", "説明", "制約"], rows=rows,
                source_sha256=hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                template="templates/deliverables/data/03-entity-definition.md",
                provenance="型はInvoiceの型注釈から抽出。PK・説明・制約はこの試験用モデルの明示メタデータ。実業務要求ではない。")
    (out / "entity.json").write_text(json.dumps(data, ensure_ascii=False, indent=2)+"\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.6)
    for name in ("Normal", "Title", "Heading 1"):
        st = doc.styles[name]
        st.font.name = "Arial Unicode MS"
        st.font.color.rgb = RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(data["title"], 0)
    doc.add_paragraph("ローカル変換試験　2026-09-26\n企業の実PJ・実顧客データを使わない合成試料。")
    doc.add_heading(data["entity"], 1)
    doc.add_paragraph(data["provenance"])
    table = doc.add_table(rows=1, cols=7)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    widths = [.3,1.35,1.0,.35,.55,2.45,1.0]
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for cell, label, width in zip(table.rows[0].cells, data["headers"], widths):
        cell.text = label
        cell.width = Inches(width)
    for row in rows:
        for cell, value, width in zip(table.add_row().cells, row, widths):
            cell.text = str(value)
            cell.width = Inches(width)
    doc.add_paragraph("空欄は元モデルの空欄を保持。導出属性はこの試料には存在しない。")
    doc.add_paragraph("参照テンプレート: "+data["template"])
    doc.add_paragraph("逆生成元: build_document.py / Invoice\nSHA256: "+data["source_sha256"])
    doc.add_paragraph("汎用ライブラリによる限定的な変換例。Plugin専用のWord/Excel変換器や任意の顧客様式への対応を示すものではない。")
    # Pin run fonts as well as styles; bundled template theme fonts may override styles.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    paragraphs = list(doc.paragraphs)
    for row in table.rows:
        for cell in row.cells:
            paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial Unicode MS"
            fonts = run._element.get_or_add_rPr().rFonts
            for attr in list(fonts.attrib):
                if "Theme" in attr:
                    del fonts.attrib[attr]
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn("w:"+key), "Arial Unicode MS")
    doc.save(out / "entity.docx")
    reopened = Document(out / "entity.docx")
    assert [[c.text for c in r.cells] for r in reopened.tables[0].rows] == [[str(c) for c in row] for row in [data["headers"], *rows]]
    print("DOCX: all table cells round-trip matched")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--output", required=True, type=Path)
    main(p.parse_args().output)
